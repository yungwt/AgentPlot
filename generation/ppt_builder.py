# generation/ppt_builder.py
import json
import os
import glob
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw
from playwright.sync_api import sync_playwright


def _calculate_textbox_position(pos_x, pos_y, width, height, alignment):
    return pos_x, pos_y


def _svg_to_png(svg_path, png_path, w_in, h_in):
    """使用 Playwright 渲染 SVG 为 PNG"""
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(
                headless=True,
                args=['--no-sandbox', '--disable-setuid-sandbox', '--disable-dev-shm-usage', '--disable-gpu']
            )
            
            width = int(w_in * 96)
            height = int(h_in * 96)
            page = browser.new_page(viewport={'width': width, 'height': height})
            
            with open(svg_path, 'r', encoding='utf-8') as f:
                svg_content = f.read()
            
            import re
            has_viewbox = re.search(r'viewBox=["\']([^"\']+)["\']', svg_content)
            has_width = re.search(r'width=["\']([0-9.]+)(?:in|px|cm)?["\']', svg_content)
            has_height = re.search(r'height=["\']([0-9.]+)(?:in|px|cm)?["\']', svg_content)
            
            if not has_viewbox and has_width and has_height:
                svg_width = float(has_width.group(1))
                svg_height = float(has_height.group(1))
                if 'in' in has_width.group(0):
                    svg_width = svg_width * 96
                    svg_height = svg_height * 96
                svg_content = svg_content.replace(
                    '<svg',
                    f'<svg viewBox="0 0 {int(svg_width)} {int(svg_height)}"'
                )
            
            html_content = f'''<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <style>
        * {{ margin: 0; padding: 0; }}
        body {{ background: white; display: flex; justify-content: center; align-items: center; width: 100vw; height: 100vh; overflow: hidden; }}
        svg {{ width: 100%; height: 100%; font-family: "Microsoft YaHei", "SimSun", "PingFang SC", sans-serif; }}
        text {{ font-family: "Microsoft YaHei", "SimSun", "PingFang SC", sans-serif; }}
    </style>
</head>
<body>{svg_content}</body>
</html>'''
            
            page.set_content(html_content, wait_until='networkidle')
            page.wait_for_timeout(500)
            page.screenshot(path=png_path, full_page=True)
            browser.close()
            return True
    except Exception as e:
        print(f"   ⚠️ SVG渲染失败: {e}")
        return False


def _placeholder_png(path, w_in, h_in, desc=""):
    w, h = int(w_in*96), int(h_in*96)
    img = Image.new('RGB', (w,h), '#F0F4F8')
    d = ImageDraw.Draw(img)
    d.rectangle([1,1,w-2,h-2], outline='#CBD5E1', width=2)
    d.text((w//2, h//2), "[ 图片占位 ]", fill='#64748B', anchor="mm")
    img.save(path, 'PNG')


def _hex_to_rgb(h):
    if not h or not h.startswith('#'):
        return RGBColor(200, 200, 200)
    h = h.lstrip('#')
    try:
        return RGBColor(*[int(h[i:i+2], 16) for i in (0, 2, 4)])
    except:
        return RGBColor(200, 200, 200)


def generate_ppt(layout_path, image_dir, output_path):
    """
    生成 PPT
    
    Args:
        layout_path: layout.json 路径
        image_dir: 图片目录（包含所有 .svg 和 .png）
        output_path: 输出 PPT 路径
    """
    with open(layout_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    
    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
    
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
    blank = prs.slide_layouts[6]
    
    temp_pngs = []  # 记录临时 PNG 文件，用于清理

    for s in data["slides"]:
        slide = prs.slides.add_slide(blank)
        page = s["page_number"]
        print(f"\n📄 第{page}页 | {s['layout_type']}")
        
        # 背景色
        try:
            bg = slide.background
            bg.fill.solid()
            bg.fill.fore_color.rgb = _hex_to_rgb(s.get("background_color", "#FAFBFC"))
        except:
            pass
        
        # 装饰
        for dec in s.get("decorations", []):
            try:
                m = {"line": MSO_SHAPE.RECTANGLE, "rect": MSO_SHAPE.RECTANGLE,
                     "circle": MSO_SHAPE.OVAL, "polygon": MSO_SHAPE.ISOSCELES_TRIANGLE,
                     "rectangle": MSO_SHAPE.RECTANGLE}
                shape = slide.shapes.add_shape(
                    m.get(dec.get("type", ""), MSO_SHAPE.RECTANGLE),
                    Inches(dec["position"]["x"]), Inches(dec["position"]["y"]),
                    Inches(dec["size"]["width"]), Inches(dec["size"]["height"])
                )
                shape.fill.solid()
                shape.fill.fore_color.rgb = _hex_to_rgb(dec.get("color", "#CCCCCC"))
                shape.line.fill.background()
            except Exception as e:
                print(f"   ⚠️ 装饰元素失败: {e}")
        
        # ===== 图片处理：统一从 image_dir 读取 =====
        if s.get("image"):
            img = s["image"]
            page = s["page_number"]
            
            # 构建图片路径
            png_path = os.path.join(image_dir, f"page_{page}.png")
            svg_path = os.path.join(image_dir, f"page_{page}.svg")
            
            img_to_insert = None
            
            # 1. 优先使用 PNG（已经生成好的）
            if os.path.exists(png_path):
                img_to_insert = png_path
                print(f"   🖼️ 使用 PNG: page_{page}.png")
            
            # 2. 如果没有 PNG，但有 SVG，临时转换为 PNG
            elif os.path.exists(svg_path):
                temp_png = f"output/temp_p{page}.png"
                temp_pngs.append(temp_png)
                if _svg_to_png(svg_path, temp_png, img["size"]["width"], img["size"]["height"]):
                    img_to_insert = temp_png
                    print(f"   🖼️ SVG 转 PNG: page_{page}.svg")
                else:
                    _placeholder_png(temp_png, img["size"]["width"], img["size"]["height"])
                    img_to_insert = temp_png
                    print(f"   🖼️ SVG 转换失败，使用占位图")
            
            # 3. 都没有，使用占位图
            else:
                placeholder = f"output/temp_placeholder_{page}.png"
                temp_pngs.append(placeholder)
                _placeholder_png(placeholder, img["size"]["width"], img["size"]["height"])
                img_to_insert = placeholder
                print(f"   🖼️ 无图片文件，使用占位图")
            
            # 插入图片
            if img_to_insert and os.path.exists(img_to_insert):
                try:
                    slide.shapes.add_picture(
                        img_to_insert,
                        Inches(img["position"]["x"]), Inches(img["position"]["y"]),
                        Inches(img["size"]["width"]), Inches(img["size"]["height"])
                    )
                except Exception as e:
                    print(f"   ⚠️ 图片插入失败: {e}")
        
        # 文字
        for block in s.get("text_blocks", []):
            try:
                left, top = _calculate_textbox_position(
                    block["position"]["x"],
                    block["position"]["y"],
                    block["size"]["width"],
                    block["size"]["height"],
                    block["style"].get("alignment", "left")
                )
                
                txBox = slide.shapes.add_textbox(
                    Inches(left), Inches(top),
                    Inches(block["size"]["width"]), Inches(block["size"]["height"])
                )
                tf = txBox.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = block["text"]
                
                style = block.get("style", {})
                if style.get("font_size"):
                    p.font.size = Pt(style["font_size"])
                if style.get("font_weight") == "bold":
                    p.font.bold = True
                if style.get("color"):
                    p.font.color.rgb = _hex_to_rgb(style["color"])
                
                amap = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
                p.alignment = amap.get(style.get("alignment", "left"), PP_ALIGN.LEFT)
                print(f"   📝 {block['text'][:40]}...")
            except Exception as e:
                print(f"   ⚠️ 文字块失败: {e}")
    
    # 保存 PPT
    prs.save(output_path)
    print(f"\n✅ PPT 已保存: {output_path}")
    
    # 清理临时 PNG
    for f in temp_pngs:
        try:
            if os.path.exists(f):
                os.remove(f)
        except:
            pass